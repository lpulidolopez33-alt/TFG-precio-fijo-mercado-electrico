from __future__ import annotations

import math
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from config import (
    RUTA_CALCULADORA_PRECIOS,
    HOJA_CALCULADORA_COMMODITY,
    CELDA_COMMODITY_FINAL,
    CELDA_COBERTURA_FINAL,
    CELDA_ETIQUETA_COMMODITY_FINAL,
    CELDA_ETIQUETA_COBERTURA_FINAL,
)


def _obtener_valor_resumen(
    resumen_decision: pd.DataFrame,
    columna: str,
) -> float:
    """Obtiene y valida un valor numérico del resumen de cobertura."""

    if resumen_decision is None or resumen_decision.empty:
        raise ValueError(
            "El resumen de decisión está vacío y no se puede exportar "
            "a la calculadora."
        )

    if columna not in resumen_decision.columns:
        raise ValueError(
            f"No existe la columna '{columna}' en el resumen de decisión."
        )

    valor = pd.to_numeric(
        resumen_decision.iloc[0][columna],
        errors="coerce",
    )

    if pd.isna(valor) or not math.isfinite(float(valor)):
        raise ValueError(
            f"El valor de la columna '{columna}' no es válido: {valor}"
        )

    return float(valor)


def _crear_copia_seguridad(ruta_excel: Path) -> Path:
    """
    Crea una copia de seguridad únicamente si todavía no existe.
    """

    ruta_backup = ruta_excel.with_name(
        f"{ruta_excel.stem}_backup_antes_exportacion{ruta_excel.suffix}"
    )

    if not ruta_backup.exists():
        shutil.copy2(ruta_excel, ruta_backup)

    return ruta_backup


def exportar_resultados_calculadora(
    resumen_decision: pd.DataFrame,
) -> None:
    """
    Exporta a la hoja COMMODITY de la calculadora:

    - la commodity final recomendada, en EUR/MWh;
    - el porcentaje total de cobertura, almacenado como decimal.

    La función modifica únicamente las celdas configuradas en config.py.
    """

    ruta_excel = Path(RUTA_CALCULADORA_PRECIOS)

    if not ruta_excel.exists():
        raise FileNotFoundError(
            "No se ha encontrado la calculadora de precios en:\n"
            f"{ruta_excel}\n"
            "Guarda el archivo junto a main.py con el nombre "
            "'Calculadora Precios.xlsx'."
        )

    commodity_final = _obtener_valor_resumen(
        resumen_decision=resumen_decision,
        columna="commodity_recomendada_EUR_MWh",
    )

    porcentaje_cubierto = _obtener_valor_resumen(
        resumen_decision=resumen_decision,
        columna="porcentaje_cobertura_total",
    )

    if not 0 <= porcentaje_cubierto <= 1:
        raise ValueError(
            "El porcentaje de cobertura debe estar comprendido entre 0 y 1. "
            f"Valor recibido: {porcentaje_cubierto}"
        )

    ruta_backup = _crear_copia_seguridad(ruta_excel)

    mantener_vba = ruta_excel.suffix.lower() == ".xlsm"

    try:
        libro = load_workbook(
            filename=ruta_excel,
            keep_vba=mantener_vba,
            keep_links=True,
        )
    except PermissionError as exc:
        raise PermissionError(
            "No se puede abrir la calculadora. Comprueba que el archivo "
            "esté cerrado en Excel."
        ) from exc

    if HOJA_CALCULADORA_COMMODITY not in libro.sheetnames:
        libro.close()
        raise ValueError(
            f"No existe la hoja '{HOJA_CALCULADORA_COMMODITY}' "
            f"en {ruta_excel.name}."
        )

    hoja = libro[HOJA_CALCULADORA_COMMODITY]

    hoja[CELDA_ETIQUETA_COMMODITY_FINAL] = (
        "Commodity final decidida [EUR/MWh]"
    )
    hoja[CELDA_COMMODITY_FINAL] = commodity_final
    hoja[CELDA_COMMODITY_FINAL].number_format = "0.00"

    hoja[CELDA_ETIQUETA_COBERTURA_FINAL] = (
        "Porcentaje cubierto final"
    )
    hoja[CELDA_COBERTURA_FINAL] = porcentaje_cubierto
    hoja[CELDA_COBERTURA_FINAL].number_format = "0.00%"

    try:
        libro.save(ruta_excel)
    except PermissionError as exc:
        libro.close()
        raise PermissionError(
            "No se puede guardar la calculadora. Comprueba que el archivo "
            "esté cerrado en Excel."
        ) from exc
    finally:
        libro.close()

    print("\nEXPORTACIÓN A LA CALCULADORA")
    print("-" * 50)
    print(f"Archivo actualizado: {ruta_excel}")
    print(f"Hoja actualizada: {HOJA_CALCULADORA_COMMODITY}")
    print(
        f"Commodity final: {commodity_final:,.2f} EUR/MWh "
        f"-> {CELDA_COMMODITY_FINAL}"
    )
    print(
        f"Porcentaje cubierto: {porcentaje_cubierto:.2%} "
        f"-> {CELDA_COBERTURA_FINAL}"
    )
    print(f"Copia de seguridad: {ruta_backup}")
    print("-" * 50)